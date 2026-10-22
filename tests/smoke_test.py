"""End-to-end smoke test: build a deck from a template using the helpers, then audit it.

Run:  python tests/smoke_test.py
"""
import sys
import tempfile
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "lib"))

from pptx import Presentation
from pptx.util import Inches, Pt

import slide_helpers as h


def make_fake_template(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "SAMPLE SLIDE FROM TEMPLATE"
    prs.save(path)


def main():
    tmp = Path(tempfile.mkdtemp())
    tpl, out = tmp / "template.pptx", tmp / "deck.pptx"
    make_fake_template(tpl)
    tpl_bytes = tpl.read_bytes()

    prs = h.open_template(tpl, out)
    layouts = h.describe_layouts(prs)
    assert layouts, "no layouts found"
    h.delete_all_slides(prs)
    assert len(prs.slides) == 0

    layout = h.get_layout(prs, "Title Only")
    s1 = prs.slides.add_slide(layout)
    s1.shapes.title.text = "Our roadmap is a journey, not a straight line"
    s1.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    h.squiggle_underline(s1, Inches(0.6), Inches(1.75), Inches(6))
    h.hand_circle(s1, Inches(8), Inches(3), Inches(2.2), Inches(1.4))
    h.hand_arrow(s1, Inches(5), Inches(3.2), Inches(2.5), Inches(1.5))
    h.journey_path(s1, Inches(0.6), Inches(5), Inches(10), Inches(1.6))
    h.set_notes(s1, "Talking point.")

    s2 = prs.slides.add_slide(layout)
    s2.shapes.title.text = "Second slide"
    h.set_notes(s2, "More notes.")

    prs.save(out)

    assert tpl.read_bytes() == tpl_bytes, "template was modified"
    try:
        h.open_template(tpl, tpl)
        raise SystemExit("open_template should refuse same path")
    except ValueError:
        pass

    # Round-trip + integrity: reopen, check XML parts, no orphaned slides from template
    re = Presentation(out)
    assert len(re.slides) == 2
    with zipfile.ZipFile(out) as z:
        assert z.testzip() is None
        slide_parts = [n for n in z.namelist() if n.startswith("ppt/slides/slide")]
        assert len(slide_parts) == 2, f"orphan slide parts: {slide_parts}"
        xml = z.read("ppt/slides/slide1.xml").decode()
        assert "custGeom" in xml and "cubicBezTo" in xml and "tailEnd" in xml

    issues = h.audit(out, template_layout_names=[l.name for l in re.slide_layouts])
    assert issues == [], issues
    print("OK - slides:", len(re.slides), "layouts:", len(layouts))


if __name__ == "__main__":
    main()
